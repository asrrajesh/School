from io import BytesIO

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm


_ASSESSMENT_CATEGORY_SHORT = {
    "fa": "FA",
    "sa": "SA",
}

_QUESTION_TYPE_LABELS = {
    "mcq": "Choose the correct answer",
    "fib": "Fill in the blanks with a suitable correct answer",
    "mtf": "Match the following",
    "tf": "Answer the following statements are True or False",
    "sa": "Answer the following in one sentence",
    "short": "Answer the following questions (Short)",
    "long": "Answer the following questions (Long)",
    "diagram": "Draw a neat labeled diagram of the following",
}

_ROMAN = ["I", "II", "III", "IV", "V", "VI", "VII", "VIII", "IX", "X"]


def _heading_for(assessment_category: str, assessment_number) -> str:
    short = _ASSESSMENT_CATEGORY_SHORT.get(
        (assessment_category or "").lower(),
        (assessment_category or "").upper(),
    )
    return f"{short} {assessment_number}".strip()


def _label_for_question_type(question_type: str) -> str:
    return _QUESTION_TYPE_LABELS.get(
        (question_type or "").lower(),
        (question_type or "").title() or "Questions",
    )


def _group_questions_by_configured_order(document: dict):
    questions = document.get("questions") or []
    configuration = document.get("configuration") or []

    ordered_types = []
    seen = set()

    for row in configuration:
        qtype = (row.get("questionType") or "").lower()

        if qtype and qtype not in seen:
            seen.add(qtype)
            ordered_types.append(qtype)

    if not ordered_types:
        for question in questions:
            qtype = (question.get("type") or "").lower()

            if qtype and qtype not in seen:
                seen.add(qtype)
                ordered_types.append(qtype)

    sections = []

    for qtype in ordered_types:
        qlist = [
            q
            for q in questions
            if (q.get("type") or "").lower() == qtype
        ]

        if qlist:
            sections.append(
                {
                    "label": _label_for_question_type(qtype),
                    "questions": qlist,
                }
            )

    return sections


def _marks_summary(document: dict):
    questions = document.get("questions") or []

    total_questions = len(questions)

    total_marks = sum(
        float(q.get("marks") or 0)
        for q in questions
    )

    configuration = document.get("configuration") or []

    distinct_marks = {
        row.get("marksPerQuestion")
        for row in configuration
        if row.get("marksPerQuestion") is not None
    }

    marks_per_question = None

    if len(distinct_marks) == 1:
        marks_per_question = next(iter(distinct_marks))

    return {
        "total_questions": total_questions,
        "total_marks": (
            int(total_marks)
            if total_marks == int(total_marks)
            else total_marks
        ),
        "marks_per_question": marks_per_question,
    }


def generate_question_paper_docx(document_data: dict) -> bytes:
    """
    Returns .docx bytes
    """

    doc = Document()

    section = doc.sections[0]

    section.top_margin = Cm(2)
    section.bottom_margin = Cm(2)
    section.left_margin = Cm(2)
    section.right_margin = Cm(2)

    heading = _heading_for(
        document_data.get("assessmentCategory"),
        document_data.get("assessmentNumber"),
    )

    marks = _marks_summary(document_data)

    sections = _group_questions_by_configured_order(document_data)

    #
    # Title
    #
    p = doc.add_paragraph()

    p.alignment = WD_ALIGN_PARAGRAPH.CENTER

    run = p.add_run(heading)
    run.bold = True

    #
    # Subject + Marks
    #
    table = doc.add_table(rows=1, cols=2)

    table.autofit = True

    row = table.rows[0]

    subj_cell = row.cells[0]
    marks_cell = row.cells[1]

    subj_cell.text = f"Subject: {document_data.get('subject', '')}"

    if marks["marks_per_question"] is not None:
        marks_text = (
            f"Marks: {marks['total_questions']} x "
            f"{marks['marks_per_question']} = "
            f"{marks['total_marks']}"
        )
    else:
        marks_text = f"Marks: {marks['total_marks']}"

    marks_cell.text = marks_text

    marks_cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT

    doc.add_paragraph()

    #
    # Questions
    #
    question_number = 1

    for idx, section_data in enumerate(sections):
        numeral = (
            _ROMAN[idx]
            if idx < len(_ROMAN)
            else str(idx + 1)
        )

        heading_para = doc.add_paragraph()

        run = heading_para.add_run(
            f"{numeral}. {section_data['label']}"
        )

        run.bold = True

        for question in section_data["questions"]:
            q_para = doc.add_paragraph()

            q_para.add_run(f"{question_number}. ").bold = True

            q_para.add_run(
                question.get("question", "")
            )

            marks_value = question.get("marks", 0)

            q_para.add_run(
                f"  [{marks_value} "
                f"{'mark' if marks_value == 1 else 'marks'}]"
            ).italic = True

            #
            # MCQ options
            #
            if (
                question.get("type") == "mcq"
                and question.get("options")
            ):
                options = question.get("options")

                labels = ["A", "B", "C", "D", "E", "F"]

                for opt_idx, option in enumerate(options):
                    opt_para = doc.add_paragraph(
                        style="List Bullet"
                    )

                    opt_para.paragraph_format.left_indent = Cm(0.75)

                    label = (
                        labels[opt_idx]
                        if opt_idx < len(labels)
                        else str(opt_idx + 1)
                    )

                    opt_para.add_run(
                        f"{label}. {option}"
                    )

            question_number += 1

    #
    # Answer Key Page
    #
    doc.add_page_break()

    answer_heading = doc.add_paragraph()

    answer_heading.alignment = WD_ALIGN_PARAGRAPH.CENTER

    answer_heading.add_run(
        "Answer Key"
    ).bold = True

    question_number = 1

    for section_data in sections:
        for question in section_data["questions"]:

            if (
                question.get("type") == "mcq"
                and question.get("correctOption")
            ):
                answer_text = question["correctOption"]

                if question.get("answerKey"):
                    answer_text += (
                        f" - {question['answerKey']}"
                    )
            else:
                answer_text = (
                    question.get("answerKey") or ""
                )

            para = doc.add_paragraph()

            para.add_run(
                f"{question_number}. "
            ).bold = True

            para.add_run(answer_text)

            question_number += 1

    buffer = BytesIO()

    doc.save(buffer)

    buffer.seek(0)

    return buffer.getvalue()